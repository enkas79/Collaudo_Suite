"""Small client for the read-only JARVIS ticket API used by MAP extraction."""

from __future__ import annotations

import json
import os
import re
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Callable
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from .core import (
    ChecklistItem,
    _coerce_map_date,
    _map_filter_matches,
    _map_status_is_rejected,
    build_ticket_url,
    map_period_bounds,
)

JARVIS_BASE_URL = "https://jarvis.breton.it"
TICKET_SEARCH_PATH = "/tickets/api/public/ticket/search?api-version=2"
TICKET_DETAIL_PATH = "/tickets/api/public/ticket/detail?api-version=2"
COMMERCIAL_CODE_FIELD = "propertydefinition_2674/matnr"


def _as_dict(value: object) -> dict:
    return value if isinstance(value, dict) else {}


def _error_text(payload: object) -> str:
    data = _as_dict(payload)
    error = data.get("error")
    if isinstance(error, dict):
        return str(error.get("message") or error.get("description") or error)
    return str(error or data.get("message") or "Risposta JARVIS non valida")


def _request_json(url: str, token: str, payload: dict, timeout: float) -> dict:
    request = Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Accept": "application/json",
            "Content-Type": "application/json",
            "jarvis-auth-token": token,
        },
        method="POST",
    )
    try:
        with urlopen(request, timeout=timeout) as response:
            raw = response.read().decode("utf-8")
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"JARVIS ha rifiutato la richiesta ({exc.code}): {detail[:300]}") from exc
    except URLError as exc:
        raise RuntimeError(f"Impossibile raggiungere JARVIS: {exc.reason}") from exc
    except TimeoutError as exc:
        raise RuntimeError("Timeout nella comunicazione con JARVIS.") from exc

    try:
        result = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError("JARVIS ha restituito una risposta non JSON.") from exc
    if not isinstance(result, dict):
        raise RuntimeError("JARVIS ha restituito una risposta inattesa.")
    if result.get("success") is False:
        raise RuntimeError(_error_text(result))
    return result


def _property_values(value: object) -> list[str]:
    """Flatten likely custom-property representations into searchable strings."""
    values: list[str] = []
    if isinstance(value, dict):
        for _key, nested in value.items():
            values.extend(_property_values(nested))
    elif isinstance(value, list):
        for nested in value:
            values.extend(_property_values(nested))
    elif value is not None:
        values.append(str(value))
    return values


def _ticket_codes(ticket: dict) -> list[str]:
    # JARVIS OmniSearch exposes the commercial code under the same inner-field
    # key used by the UI filter. Never fall back to arbitrary property values:
    # those include opaque references such as DatasetElement_706033.
    for prop in ticket.get("stringProperties", []) if isinstance(ticket.get("stringProperties"), list) else []:
        if not isinstance(prop, dict) or str(prop.get("key", "")).casefold() != COMMERCIAL_CODE_FIELD.casefold():
            continue
        value = prop.get("value")
        values = value if isinstance(value, list) else [value]
        return [str(item).strip() for item in values if item not in (None, "")]

    properties = ticket.get("properties")
    if not isinstance(properties, dict):
        return []
    named_candidates: list[str] = []
    for key, value in properties.items():
        key_text = str(key).casefold()
        if key_text == COMMERCIAL_CODE_FIELD.casefold() or any(
            term in key_text for term in ("commercial", "codice commerciale", "commessa s codice")
        ):
            named_candidates.extend(_property_values(value))
    return named_candidates


def _matches_commercial_code(codes: list[str], title: str, number: str, filter_text: str) -> bool:
    return any(_map_filter_matches(code, title, number, filter_text) for code in codes)


def _select_commercial_code(codes: list[str]) -> str:
    for code in codes:
        value = str(code).strip()
        if re.match(r"(?i)^(nc300|genya|trinity)", value):
            return value
    return str(codes[0]).strip() if codes else ""


def _excel_cell_value(value: object) -> object:
    """Convert JSON arrays/objects to readable Excel cell text."""
    if isinstance(value, list):
        return ", ".join(str(_excel_cell_value(item)) for item in value)
    if isinstance(value, dict):
        if value.get("name") is not None:
            return str(value["name"])
        if value.get("value") is not None:
            return _excel_cell_value(value["value"])
        return "; ".join(f"{key}: {_excel_cell_value(item)}" for key, item in value.items())
    return "" if value is None else value


def _ticket_number(ticket: dict) -> str:
    value = ticket.get("number", ticket.get("id", ""))
    if isinstance(value, str):
        match = re.fullmatch(r"Ticket_(\d+)", value.strip(), re.IGNORECASE)
        if match:
            value = match.group(1)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def _search_page(token: str, payload: dict, timeout: float) -> dict:
    return _request_json(JARVIS_BASE_URL + TICKET_SEARCH_PATH, token, payload, timeout)


def _fetch_ticket_rows(
    token: str,
    *,
    page_size: int,
    max_pages: int,
    timeout: float,
    progress_callback: Callable[[int, int], None] | None = None,
) -> list[dict]:
    """Fetch tickets using JARVIS' native Commercial code filter."""
    rows: list[dict] = []
    size = max(1, min(int(page_size), 500))
    for page in range(max(1, int(max_pages))):
        payload = {
            "chunkSize": size,
            "startIndex": page * size,
            "showArchived": False,
            "text": "",
            "includeAllTerms": None,
            "requestType": 0,
            # Public API form of the commercial-code filter seen in the HAR.
            # An empty contains value selects the field without restricting
            # the download to one specific machine code.
            "filters": [{
                "field": COMMERCIAL_CODE_FIELD,
                "dataType": 0,
                "value": "",
                "values": [],
                "exactSearch": False,
                "caseSensitiveSearch": False,
            }],
        }
        response = _search_page(token, payload, timeout)
        result = _as_dict(response.get("data"))
        page_rows = result.get("items", response.get("items"))
        if not isinstance(page_rows, list) or not page_rows:
            break
        total_value = result.get("total", response.get("total"))
        total = total_value if isinstance(total_value, int) else 0
        for raw in page_rows:
            raw_item = _as_dict(raw)
            result_item = raw_item.get("item", raw_item)
            ticket = dict(_as_dict(result_item))
            ticket["id"] = _ticket_number({"id": ticket.get("id", "")})
            ticket["number"] = ticket["id"]
            ticket["properties"] = {
                COMMERCIAL_CODE_FIELD: _ticket_codes(ticket),
            }
            # OmniSearch stores ticket fields in typed property arrays. Map the
            # fields used by the MAP cache while preserving the exact code.
            string_properties = ticket.get("stringProperties", [])
            values_by_key = {
                str(prop.get("key", "")).casefold(): prop.get("value", [])
                for prop in string_properties
                if isinstance(prop, dict)
            } if isinstance(string_properties, list) else {}
            def first_value(key: str, default: str = "") -> str:
                value = values_by_key.get(key.casefold(), [])
                if isinstance(value, list):
                    value = value[0] if value else default
                return str(value or default).strip()
            ticket["model"] = first_value("ticketmodelname", "Segnalazione MAP")
            ticket["title"] = str(ticket.get("title") or first_value("joip_title"))
            ticket["businessLine"] = first_value("propertydefinition_1288/jarvisformfield_3168")
            ticket["status"] = first_value("statuslabel")
            ticket["createdAt"] = ticket.get("lastUpdated") or first_value("propertydefinition_2667/last_modify_date")
            ticket["lastModifyDate"] = ticket["createdAt"]
            ticket["assignee"] = first_value("assignedtoname")
            ticket["createdBy"] = first_value("createdbyname")
            ticket["lastModifyUser"] = first_value("lastchangeuser")
            ticket["repliesNumber"] = first_value("replies_no")
            ticket["attachmentsNumber"] = first_value("attacchments_no")
            rows.append(ticket)
            processed = len(rows)
            percent = int(processed * 100 / total) if total else 0
            if progress_callback:
                progress_callback(min(99, percent), processed)
        total = result.get("total", response.get("total"))
        if len(page_rows) < size or (isinstance(total, int) and page * size + len(page_rows) >= total):
            break
    return rows


def sync_map_xlsx_from_jarvis(
    token: str,
    output_path: str | Path,
    *,
    page_size: int = 100,
    max_pages: int = 20,
    timeout: float = 20.0,
    progress_callback: Callable[[int, int], None] | None = None,
) -> Path:
    """Create a local MAP-compatible Excel cache from JARVIS tickets."""
    token = str(token or "").strip()
    if not token:
        raise ValueError("Inserire un token JARVIS con permesso ticket.read.")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    rows = _fetch_ticket_rows(
        token,
        page_size=page_size,
        max_pages=max_pages,
        timeout=timeout,
        progress_callback=progress_callback,
    )

    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("La dipendenza openpyxl è necessaria per creare la cache MAP.") from exc

    headers = [
        "Numero Ticket", "Modello ticket", "Titolo", "Linea di business", "Stato",
        "Commessa S\nCodice commerciale", "Data creazione", "🔔", "Data scadenza",
        "Data chiusura", "Ore Non Qualità", "Ore evasione attività", "In carico a",
        "Destinatari", "Visibile A", "Tags", "Creato da", "N. risposte",
        "N. allegati", "Autore ultima modifica", "Data ultima modifica",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(headers)
    existing_rows: dict[str, tuple[list[object], str]] = {}
    if output.exists():
        from openpyxl import load_workbook

        old_workbook = load_workbook(output, read_only=False, data_only=False)
        try:
            old_sheet = old_workbook["Data"] if "Data" in old_workbook.sheetnames else old_workbook.active
            for row_number in range(2, old_sheet.max_row + 1):
                values = [old_sheet.cell(row=row_number, column=column).value for column in range(1, len(headers) + 1)]
                old_ticket = str(values[0] or "").strip()
                old_code = str(values[5] or "").strip()
                if old_ticket and old_code and not re.match(r"(?i)^(DatasetElement|UserDefinedValueListEntry)_\d+$", old_code):
                    old_link = old_sheet.cell(row=row_number, column=1).hyperlink
                    existing_rows[old_ticket] = (values, old_link.target if old_link else "")
        finally:
            old_workbook.close()
    current_tickets: set[str] = set()
    for ticket in rows:
        codes = _ticket_codes(ticket)
        ticket_number = _ticket_number(ticket)
        current_tickets.add(ticket_number)
        sheet.append([
            ticket_number,
            _excel_cell_value(ticket.get("model") or ticket.get("modelId") or ""),
            _excel_cell_value(ticket.get("title") or ticket.get("subject") or ""),
            _excel_cell_value(ticket.get("businessLine") or ticket.get("lineOfBusiness") or ""),
            _excel_cell_value(ticket.get("status") or ticket.get("statusLabel") or ""),
            _select_commercial_code(codes),
            _excel_cell_value(ticket.get("createdAt") or ""),
            "",
            _excel_cell_value(ticket.get("dueDate") or ""),
            _excel_cell_value(ticket.get("closingDate") or ""),
            _excel_cell_value(ticket.get("nonQualityHours") or ""),
            _excel_cell_value(ticket.get("workHours") or ""),
            _excel_cell_value(ticket.get("assignee") or ""),
            _excel_cell_value(ticket.get("recipients") or ""),
            _excel_cell_value(ticket.get("visibleTo") or ""),
            _excel_cell_value(ticket.get("tags") or ""),
            _excel_cell_value(ticket.get("createdBy") or ""),
            _excel_cell_value(ticket.get("repliesNumber") or ""),
            _excel_cell_value(ticket.get("attachmentsNumber") or ""),
            _excel_cell_value(ticket.get("lastModifyUser") or ""),
            _excel_cell_value(ticket.get("lastModifyDate") or ""),
        ])
        if ticket_number:
            ticket_cell = sheet.cell(row=sheet.max_row, column=1)
            ticket_cell.hyperlink = build_ticket_url(ticket_number)
            ticket_cell.style = "Hyperlink"
    # Preserve older cache entries that are outside the current API page window.
    for old_ticket, (values, old_link) in existing_rows.items():
        if old_ticket in current_tickets:
            continue
        sheet.append(values)
        if old_link:
            old_cell = sheet.cell(row=sheet.max_row, column=1)
            old_cell.hyperlink = old_link
            old_cell.style = "Hyperlink"
    temporary_fd, temporary_name = tempfile.mkstemp(
        prefix=f"{output.stem}_",
        suffix=".tmp.xlsx",
        dir=output.parent,
    )
    os.close(temporary_fd)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        try:
            os.replace(temporary, output)
            saved_output = output
        except PermissionError:
            # Windows cannot replace a workbook while another process holds it
            # open (often Excel or an antivirus scanner). Keep the old cache intact
            # and install this complete generation under a fresh sibling filename.
            suffix = datetime.now().strftime("%Y%m%d_%H%M%S")
            saved_output = output.with_name(f"{output.stem}_{suffix}{output.suffix}")
            sequence = 1
            while saved_output.exists():
                saved_output = output.with_name(f"{output.stem}_{suffix}_{sequence}{output.suffix}")
                sequence += 1
            os.replace(temporary, saved_output)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
    if progress_callback:
        progress_callback(100, len(rows))
    return saved_output


def load_map_items_from_jarvis(
    token: str,
    filter_text: str = "",
    *,
    reference_date: date | datetime | str | None = None,
    months_back: int | None = None,
    page_size: int = 100,
    max_pages: int = 20,
    timeout: float = 20.0,
) -> list[ChecklistItem]:
    """Fetch ticket MAP candidates from JARVIS and apply the existing MAP rules locally.

    The API schema does not define a standard Commercial code field, so the function
    reads common custom-property names from ticket details when a filter is supplied.
    """
    token = str(token or "").strip()
    if not token:
        raise ValueError("Inserire un token JARVIS con permesso ticket.read.")
    period_start, period_end = (None, None)
    if months_back is not None:
        period_start, period_end = map_period_bounds(reference_date, int(months_back))

    result: list[ChecklistItem] = []
    seen: set[str] = set()
    normalized_filter = str(filter_text or "").strip()
    for page in range(max(1, int(max_pages))):
        payload = {
            "pageSize": max(1, min(int(page_size), 500)),
            "startIndex": page * max(1, min(int(page_size), 500)),
            "showArchived": False,
            "sortField": "lastModifyDate",
            "sortDescending": True,
        }
        # Do not send the Commercial code as full-text query: in some JARVIS
        # installations it is a custom property and is not indexed in Text.
        # Filtering is applied after reading ticket properties below.
        response = _search_page(token, payload, timeout)
        rows = response.get("data")
        if not isinstance(rows, list) or not rows:
            break
        for raw in rows:
            ticket = _as_dict(raw)
            if _map_status_is_rejected(ticket.get("status") or ticket.get("statusLabel")):
                continue
            title = str(ticket.get("title") or ticket.get("subject") or "").strip()
            number = _ticket_number(ticket)
            if not title:
                continue
            raw_date = ticket.get("lastModifyDate") or ticket.get("createdAt")
            ticket_date = _coerce_map_date(raw_date)
            if period_start and (ticket_date is None or ticket_date < period_start or ticket_date > period_end):
                continue
            codes = _ticket_codes(ticket)
            if normalized_filter and not _matches_commercial_code(codes, title, number, normalized_filter):
                ticket_id = str(ticket.get("id") or "").strip()
                if ticket_id:
                    detail = _request_json(
                        JARVIS_BASE_URL + TICKET_DETAIL_PATH,
                        token,
                        {"ticketId": ticket_id},
                        timeout,
                    ).get("data")
                    detail_dict = _as_dict(detail)
                    codes = _ticket_codes(detail_dict)
                if not _matches_commercial_code(codes, title, number, normalized_filter):
                    continue
            key = (number + " " + title).casefold().strip()
            if not key or key in seen:
                continue
            seen.add(key)
            result.append(
                ChecklistItem(
                    text=title,
                    source="JARVIS API",
                    kind="",
                    ticket_number=number,
                    ticket_url=build_ticket_url(number),
                )
            )
        total = response.get("total")
        if len(rows) < payload["pageSize"] or (isinstance(total, int) and payload["startIndex"] + len(rows) >= total):
            break
    return result
