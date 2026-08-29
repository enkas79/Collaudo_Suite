import tempfile
from datetime import date, datetime
import unittest
from pathlib import Path

from openpyxl import Workbook

from collaudo_suite.analyzer.analysis import ExcelAnalyzer
from collaudo_suite.analyzer.excel_utils import cell_to_indices, col_to_index, index_to_col
from collaudo_suite.analyzer.models import AnalysisParams
from collaudo_suite.analyzer.similarity import UnionFind, calculate_similarity
from collaudo_suite.analyzer.text_utils import TextNormalizer
from collaudo_suite.checklist.core import (
    get_default_fixed_docx_path,
    get_default_map_xlsx_path,
    load_default_fixed_items,
    load_map_items_for_filter,
    ticket_desc_sort_key,
    map_period_bounds,
)
from collaudo_suite.control_exchange import ExternalControl, export_controls_json, export_controls_xlsx, import_controls


class ExcelUtilsTests(unittest.TestCase):
    def test_col_conversion(self):
        self.assertEqual(col_to_index("A"), 0)
        self.assertEqual(col_to_index("AA"), 26)
        self.assertEqual(index_to_col(26), "AA")
        self.assertEqual(cell_to_indices("B14"), (13, 1))
        self.assertEqual(cell_to_indices("14"), (13, -1))


class TextNormalizerTests(unittest.TestCase):
    def test_preserves_technical_codes(self):
        norm = TextNormalizer(stemming=False).normalize("Allarme E32.0 su DM30 e filo 30F-6")
        self.assertIn("E32.0", norm.technical_codes)
        self.assertIn("DM30", norm.technical_codes)
        self.assertIn("30F-6", norm.technical_codes)


class SimilarityTests(unittest.TestCase):
    def test_union_find_connected_component(self):
        uf = UnionFind(4)
        uf.union(0, 1)
        uf.union(1, 2)
        groups = [sorted(v) for v in uf.groups().values()]
        self.assertIn([0, 1, 2], groups)

    def test_similarity_score(self):
        score = calculate_similarity(
            "allarme pressostato aria",
            "pressostato aria generale",
            {"allarme", "pressostato", "aria"},
            {"pressostato", "aria", "generale"},
            "combinato",
        )
        self.assertGreater(score, 40)


class ExchangeTests(unittest.TestCase):
    def test_export_import_roundtrip(self):
        controls = [
            ExternalControl(
                control="Verificare il sensore di livello minimo",
                original="Sensore livello minimo non intervenuto",
                machine="GENYA",
                category="Sensori",
                ticket="12345",
                occurrences=8,
                latest_date="2026-07-08",
            )
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = export_controls_xlsx(controls, Path(tmp) / "controlli.xlsx")
            loaded = import_controls(path)
        self.assertEqual(len(loaded), 1)
        self.assertEqual(loaded[0].control, controls[0].control)
        self.assertEqual(loaded[0].ticket, "12345")
        self.assertEqual(loaded[0].occurrences, 8)

    def test_csv_and_json_import(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            csv_path = tmp_path / "controlli.csv"
            csv_path.write_text(
                "CONTROLLO;MACCHINA;TICKET;OCCORRENZE\n"
                "Verificare il pressostato;TRINITY;9876;4\n",
                encoding="utf-8-sig",
            )
            csv_loaded = import_controls(csv_path)
            self.assertEqual(csv_loaded[0].machine, "TRINITY")
            self.assertEqual(csv_loaded[0].ticket, "9876")

            json_path = export_controls_json(csv_loaded, tmp_path / "controlli.json")
            json_loaded = import_controls(json_path)
            self.assertEqual(json_loaded[0].control, "Verificare il pressostato")
            self.assertEqual(json_loaded[0].occurrences, 4)

    def test_import_analyzer_summary_as_controls(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report_analyzer.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sintesi"
            sheet.append(["CONCETTO PRINCIPALE DEL GRUPPO", "NUMERO DI OCCORRENZE", "DATA PIU RECENTE"])
            sheet.append(["Sensore non intervenuto", 6, "2026-07-08"])
            workbook.save(path)
            loaded = import_controls(path)
        self.assertEqual(loaded[0].control, "Sensore non intervenuto")
        self.assertEqual(loaded[0].occurrences, 6)


class ChecklistResourcesTests(unittest.TestCase):
    def test_bundled_resources_load(self):
        self.assertTrue(get_default_fixed_docx_path().exists())
        self.assertGreater(len(load_default_fixed_items()), 0)

        map_path = get_default_map_xlsx_path()
        if map_path.exists():
            self.assertGreater(len(load_map_items_for_filter(map_path, "")), 0)


class AnalyzerEndToEndTests(unittest.TestCase):
    def test_groups_similar_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "anomalie.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Punti Aperti"
            for row in range(1, 14):
                sheet.cell(row=row, column=2, value="")
            values = [
                "Mancato intervento sensore livello minimo",
                "Sensore di livello minimo non intervenuto",
                "Il sensore livello minimo non interviene",
                "Perdita raccordo pneumatico",
            ]
            for offset, value in enumerate(values, start=14):
                sheet.cell(row=offset, column=2, value=value)
                sheet.cell(row=offset, column=5, value=f"2026-07-{offset-10:02d}")
            workbook.save(path)

            report = ExcelAnalyzer().run(
                AnalysisParams(
                    files=[str(path)],
                    threshold=45,
                    min_occurrences=2,
                    start_cell="B14",
                    target_col="B",
                    date_col="E",
                    sheet_name="Punti Aperti",
                    stemming=False,
                )
            )
        self.assertTrue(report.summary)
        self.assertGreaterEqual(report.summary[0]["NUMERO DI OCCORRENZE"], 2)


if __name__ == "__main__":
    unittest.main()

class MapTicketLinkTests(unittest.TestCase):
    def test_reads_ticket_hyperlink(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Commercial code", "Title", "Ticket Number"])
            sheet.append(["IGenya", "Verificare anomalia", "12345"])
            sheet["C2"].hyperlink = "https://example.invalid/ticket/12345"
            workbook.save(path)
            items = load_map_items_for_filter(path, "IGenya")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].ticket_number, "12345")
        self.assertEqual(items[0].ticket_url, "https://example.invalid/ticket/12345")

class MapTicketGeneratedLinkTests(unittest.TestCase):
    def test_builds_direct_ticket_link_from_ticket_number(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map_without_hyperlink.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Commercial code", "Title", "Ticket Number"])
            sheet.append(["IGenya", "Verificare anomalia", "12345"])
            workbook.save(path)
            items = load_map_items_for_filter(path, "IGenya")
        self.assertEqual(len(items), 1)
        self.assertEqual(
            items[0].ticket_url,
            "https://jarvis.breton.it/UI/#/ticketing/ticket/Ticket_12345",
        )


class TicketDescendingSortTests(unittest.TestCase):
    def test_numeric_ticket_numbers_are_sorted_descending(self):
        tickets = ["99", "Ticket_100", "7", "", "ABC"]
        ordered = sorted(tickets, key=ticket_desc_sort_key, reverse=True)
        self.assertEqual(ordered, ["Ticket_100", "99", "7", "ABC", ""])


class MapTemporalFilterTests(unittest.TestCase):
    @staticmethod
    def _build_map(path: Path, rows, header="Last Modify") -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Commercial code", "Title", "Ticket Number", header])
        for row in rows:
            sheet.append(row)
        workbook.save(path)

    def test_filters_against_collaudo_date_inclusive(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map_dates.xlsx"
            self._build_map(
                path,
                [
                    ["IGenya", "Dentro periodo", "300", datetime(2026, 7, 1, 10, 30)],
                    ["IGenya", "Sul limite", "299", datetime(2026, 6, 24, 0, 0)],
                    ["IGenya", "Troppo vecchio", "298", datetime(2026, 6, 23, 23, 59)],
                    ["IGenya", "Successivo al collaudo", "301", datetime(2026, 7, 25, 0, 0)],
                ],
            )
            items = load_map_items_for_filter(
                path,
                "IGenya",
                reference_date=date(2026, 7, 24),
                months_back=1,
            )
        self.assertEqual([item.ticket_number for item in items], ["300", "299"])

    def test_accepts_italian_string_date_header(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map_dates_it.xlsx"
            self._build_map(
                path,
                [
                    ["GENYA", "Recente", "10", "15/04/2026"],
                    ["GENYA", "Vecchio", "9", "14/04/2026"],
                ],
                header="Data ultima modifica",
            )
            items = load_map_items_for_filter(
                path,
                "GENYA",
                reference_date="15/07/2026",
                months_back=3,
            )
        self.assertEqual([item.ticket_number for item in items], ["10"])

    def test_calendar_month_clamps_end_of_month(self):
        start, end = map_period_bounds(date(2026, 3, 31), 1)
        self.assertEqual(start, date(2026, 2, 28))
        self.assertEqual(end, date(2026, 3, 31))


    def test_recognized_date_column_without_valid_dates_is_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map_invalid_dates.xlsx"
            self._build_map(
                path,
                [["IGenya", "Anomalia", "123", "non disponibile"]],
            )
            with self.assertRaisesRegex(ValueError, "non contiene date valide"):
                load_map_items_for_filter(
                    path,
                    "IGenya",
                    reference_date=date(2026, 7, 24),
                    months_back=1,
                )

    def test_generic_data_does_not_match_data_module(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map_data_module.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Commercial code", "Title", "Ticket Number", "Data Module"])
            sheet.append(["IGenya", "Anomalia", "123", "DM30"])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "colonna data"):
                load_map_items_for_filter(
                    path,
                    "IGenya",
                    reference_date=date(2026, 7, 24),
                    months_back=1,
                )

    def test_missing_date_column_is_explicit_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map_no_date.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Commercial code", "Title", "Ticket Number"])
            sheet.append(["IGenya", "Anomalia", "123"])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "colonna data"):
                load_map_items_for_filter(
                    path,
                    "IGenya",
                    reference_date=date(2026, 7, 24),
                    months_back=1,
                )


class MapCacheInvalidationTests(unittest.TestCase):
    @staticmethod
    def _write_map(path: Path, title: str, ticket: str) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Commercial code", "Title", "Ticket Number"])
        sheet.append(["IGenya", title, ticket])
        workbook.save(path)

    def test_reloads_map_when_same_file_changes(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "map.xlsx"
            self._write_map(path, "Prima anomalia", "100")
            first = load_map_items_for_filter(path, "IGenya")

            self._write_map(path, "Seconda anomalia aggiornata", "101")
            second = load_map_items_for_filter(path, "IGenya")

        self.assertEqual([item.text for item in first], ["Prima anomalia"])
        self.assertEqual([item.text for item in second], ["Seconda anomalia aggiornata"])
