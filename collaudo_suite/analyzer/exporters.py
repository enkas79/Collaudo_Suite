from __future__ import annotations

from pathlib import Path
from typing import Iterable

import pandas as pd

from .excel_utils import force_extension
from .models import AnalysisReport


def extension_from_filter(selected_filter: str) -> str:
    text = (selected_filter or "").lower()
    if "word" in text or "docx" in text:
        return ".docx"
    if "text" in text or "txt" in text:
        return ".txt"
    if "csv" in text:
        return ".csv"
    return ".xlsx"


def prepare_export_path(path: str, selected_filter: str) -> str:
    return force_extension(path, extension_from_filter(selected_filter))


def _metadata_rows(report: AnalysisReport) -> list[dict[str, object]]:
    rows = [{"CHIAVE": key, "VALORE": value} for key, value in report.metadata.items()]
    for warning in report.warnings:
        rows.append({"CHIAVE": "WARNING", "VALORE": warning})
    for error in report.errors:
        rows.append({"CHIAVE": "ERROR", "VALORE": error})
    return rows


def export_report(report: AnalysisReport, path: str, selected_filter: str) -> list[str]:
    path = prepare_export_path(path, selected_filter)
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return _export_xlsx(report, path)
    if suffix == ".csv":
        return _export_csv(report, path)
    if suffix == ".txt":
        return _export_txt(report, path)
    if suffix == ".docx":
        return _export_docx(report, path)
    raise ValueError(f"Formato non supportato: {suffix}")


def _export_xlsx(report: AnalysisReport, path: str) -> list[str]:
    summary_df = pd.DataFrame(report.summary)
    detail_df = pd.DataFrame(report.detail)
    metadata_df = pd.DataFrame(_metadata_rows(report))

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if not summary_df.empty:
            summary_df.to_excel(writer, sheet_name="Sintesi", index=False)
        else:
            pd.DataFrame([{"MESSAGGIO": "Nessun risultato"}]).to_excel(writer, sheet_name="Sintesi", index=False)
        if not detail_df.empty:
            detail_df.to_excel(writer, sheet_name="Dettaglio", index=False)
        metadata_df.to_excel(writer, sheet_name="Metadati", index=False)

        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            for col_cells in sheet.columns:
                max_len = 0
                col_letter = col_cells[0].column_letter
                for cell in col_cells[:200]:
                    max_len = max(max_len, len(str(cell.value or "")))
                sheet.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 70)
                for cell in col_cells:
                    cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
    return [path]


def _export_csv(report: AnalysisReport, path: str) -> list[str]:
    base = Path(path)
    saved: list[str] = []
    pd.DataFrame(report.summary).to_csv(base, index=False, encoding="utf-8-sig")
    saved.append(str(base))

    if report.detail:
        detail_path = base.with_name(base.stem + "_dettaglio.csv")
        pd.DataFrame(report.detail).to_csv(detail_path, index=False, encoding="utf-8-sig")
        saved.append(str(detail_path))

    metadata_path = base.with_name(base.stem + "_metadati.csv")
    pd.DataFrame(_metadata_rows(report)).to_csv(metadata_path, index=False, encoding="utf-8-sig")
    saved.append(str(metadata_path))
    return saved


def _export_txt(report: AnalysisReport, path: str) -> list[str]:
    with open(path, "w", encoding="utf-8") as handle:
        handle.write("REPORT ANALISI\n")
        handle.write("=" * 80 + "\n\n")
        handle.write("METADATI\n")
        for key, value in report.metadata.items():
            handle.write(f"- {key}: {value}\n")
        if report.warnings:
            handle.write("\nAVVISI\n")
            for warning in report.warnings:
                handle.write(f"- {warning}\n")
        if report.errors:
            handle.write("\nERRORI\n")
            for error in report.errors:
                handle.write(f"- {error}\n")

        handle.write("\nSINTESI\n")
        handle.write("-" * 80 + "\n")
        if not report.summary:
            handle.write("Nessun risultato.\n")
        for row in report.summary:
            for key, value in row.items():
                handle.write(f"{key}: {value}\n")
            handle.write("-" * 80 + "\n")

        if report.detail:
            handle.write("\nDETTAGLIO\n")
            handle.write("-" * 80 + "\n")
            for row in report.detail:
                handle.write(" | ".join(f"{key}: {value}" for key, value in row.items()))
                handle.write("\n")
    return [path]


def _export_docx(report: AnalysisReport, path: str) -> list[str]:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RuntimeError("Per esportare in Word installa python-docx.") from exc

    doc = Document()
    doc.add_heading("Report Analisi", level=0)

    doc.add_heading("Metadati", level=1)
    for key, value in report.metadata.items():
        doc.add_paragraph(f"{key}: {value}")

    if report.warnings:
        doc.add_heading("Avvisi", level=1)
        for warning in report.warnings:
            doc.add_paragraph(warning, style="List Bullet")

    doc.add_heading("Sintesi", level=1)
    _add_table(doc, report.summary)

    if report.detail:
        doc.add_heading("Dettaglio", level=1)
        _add_table(doc, report.detail)

    doc.save(path)
    return [path]


def _add_table(doc, rows: list[dict[str, object]]) -> None:
    if not rows:
        doc.add_paragraph("Nessun risultato.")
        return
    columns = list(rows[0].keys())
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for idx, column in enumerate(columns):
        table.rows[0].cells[idx].text = str(column)
    for row in rows:
        cells = table.add_row().cells
        for idx, column in enumerate(columns):
            cells[idx].text = str(row.get(column, ""))
