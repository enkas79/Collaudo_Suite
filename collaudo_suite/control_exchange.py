from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass(slots=True)
class ExternalControl:
    control: str
    original: str = ""
    machine: str = ""
    category: str = ""
    ticket: str = ""
    occurrences: int = 1
    latest_date: str = ""
    source: str = "Analyzer"

    def normalized_key(self) -> str:
        return re.sub(r"\W+", "", self.control.casefold())

    def source_label(self) -> str:
        parts = [self.source.strip() or "Analyzer"]
        if self.machine.strip():
            parts.append(self.machine.strip())
        if self.category.strip():
            parts.append(self.category.strip())
        if self.occurrences > 1:
            parts.append(f"{self.occurrences} occorrenze")
        return " | ".join(parts)


STANDARD_COLUMNS = [
    "CONTROLLO",
    "ANOMALIA_ORIGINALE",
    "MACCHINA",
    "CATEGORIA",
    "TICKET",
    "OCCORRENZE",
    "DATA_PIU_RECENTE",
    "ORIGINE",
]

ALIASES = {
    "control": {
        "controllo", "control", "test", "descrizione controllo", "controllo da eseguire",
        "title", "titolo", "azione", "azione di controllo",
    },
    "original": {
        "anomalia originale", "anomalia", "originale", "descrizione anomalia",
        "concetto principale del gruppo", "testo trovato", "testo originale",
    },
    "machine": {"macchina", "machine", "tipo macchina", "commercial code", "codice commerciale"},
    "category": {"categoria", "category", "gruppo", "tipologia", "famiglia"},
    "ticket": {"ticket", "ticket number", "numero ticket", "n ticket", "nr ticket"},
    "occurrences": {"occorrenze", "numero di occorrenze", "frequency", "frequenza", "n occorrenze"},
    "latest_date": {"data piu recente", "data più recente", "latest date", "data"},
    "source": {"origine", "source", "sorgente"},
}


def _norm_header(value: object) -> str:
    text = str(value or "").strip().casefold()
    text = text.replace("_", " ")
    text = re.sub(r"\s+", " ", text)
    return text


def _field_for_header(header: object) -> str | None:
    normalized = _norm_header(header)
    for field, aliases in ALIASES.items():
        if normalized in aliases:
            return field
    return None


def _to_int(value: object, default: int = 1) -> int:
    try:
        if value is None or str(value).strip() == "":
            return default
        return max(1, int(float(str(value).replace(",", "."))))
    except (TypeError, ValueError):
        return default


def _control_from_mapping(mapping: dict[str, object]) -> ExternalControl | None:
    control = str(mapping.get("control", "") or "").strip()
    original = str(mapping.get("original", "") or "").strip()
    if not control:
        control = original
    if not control:
        return None
    return ExternalControl(
        control=control,
        original=original,
        machine=str(mapping.get("machine", "") or "").strip(),
        category=str(mapping.get("category", "") or "").strip(),
        ticket=str(mapping.get("ticket", "") or "").strip(),
        occurrences=_to_int(mapping.get("occurrences", 1)),
        latest_date=str(mapping.get("latest_date", "") or "").strip(),
        source=str(mapping.get("source", "Analyzer") or "Analyzer").strip() or "Analyzer",
    )


def deduplicate_controls(controls: Iterable[ExternalControl]) -> list[ExternalControl]:
    result: list[ExternalControl] = []
    seen: set[str] = set()
    for control in controls:
        key = control.normalized_key()
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(control)
    return result


def export_controls_xlsx(controls: Sequence[ExternalControl], path: str | Path) -> Path:
    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Controlli"
    sheet.append(STANDARD_COLUMNS)

    for control in controls:
        sheet.append([
            control.control,
            control.original,
            control.machine,
            control.category,
            control.ticket,
            control.occurrences,
            control.latest_date,
            control.source,
        ])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [60, 60, 20, 24, 18, 14, 20, 24]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    instructions = workbook.create_sheet("Istruzioni")
    instructions["A1"] = "Formato di scambio controlli - Collaudo Suite"
    instructions["A1"].font = Font(bold=True, size=14)
    instructions["A3"] = "La colonna CONTROLLO è l'unica obbligatoria. Le altre colonne servono per tracciabilità e filtri."
    instructions["A5"] = "Il file può essere importato dalla sezione Checklist > Controlli Analyzer."
    instructions.column_dimensions["A"].width = 110
    instructions["A3"].alignment = Alignment(wrap_text=True)
    instructions["A5"].alignment = Alignment(wrap_text=True)

    workbook.save(target)
    return target


def export_controls_json(controls: Sequence[ExternalControl], path: str | Path) -> Path:
    target = Path(path)
    if target.suffix.lower() != ".json":
        target = target.with_suffix(".json")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps({"format": "collaudo-suite-controls-1", "controls": [asdict(c) for c in controls]}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return target


def import_controls(path: str | Path) -> list[ExternalControl]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _import_xlsx(source)
    if suffix == ".csv":
        return _import_csv(source)
    if suffix == ".json":
        return _import_json(source)
    raise ValueError("Formato non supportato. Usa Excel .xlsx, CSV o JSON.")


def _import_xlsx(path: Path) -> list[ExternalControl]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Controlli"] if "Controlli" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    return _controls_from_rows(rows)


def _import_csv(path: Path) -> list[ExternalControl]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            rows = list(csv.reader(handle, dialect))
        except csv.Error:
            rows = list(csv.reader(handle, delimiter=";"))
    return _controls_from_rows(rows)


def _import_json(path: Path) -> list[ExternalControl]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(raw, dict):
        rows = raw.get("controls", raw.get("items", []))
    else:
        rows = raw
    controls: list[ExternalControl] = []
    if not isinstance(rows, list):
        raise ValueError("Il file JSON non contiene una lista di controlli.")
    for row in rows:
        if not isinstance(row, dict):
            continue
        mapping: dict[str, object] = {}
        for key, value in row.items():
            field = _field_for_header(key) or key
            mapping[field] = value
        control = _control_from_mapping(mapping)
        if control:
            controls.append(control)
    return deduplicate_controls(controls)


def _controls_from_rows(rows: Sequence[Sequence[object]]) -> list[ExternalControl]:
    if not rows:
        return []

    header_index = None
    field_by_col: dict[int, str] = {}
    for idx, row in enumerate(rows[:30]):
        candidate: dict[int, str] = {}
        for col, value in enumerate(row):
            field = _field_for_header(value)
            if field:
                candidate[col] = field
        if "control" in candidate.values() or "original" in candidate.values():
            header_index = idx
            field_by_col = candidate
            break

    if header_index is None:
        # Fallback: first non-empty column is interpreted as control text.
        controls = []
        for row in rows:
            value = next((str(v).strip() for v in row if v is not None and str(v).strip()), "")
            if value:
                controls.append(ExternalControl(control=value, source=path_source_label("Importazione")))
        return deduplicate_controls(controls)

    controls: list[ExternalControl] = []
    for row in rows[header_index + 1:]:
        mapping: dict[str, object] = {}
        for col, field in field_by_col.items():
            mapping[field] = row[col] if col < len(row) else ""
        control = _control_from_mapping(mapping)
        if control:
            controls.append(control)
    return deduplicate_controls(controls)


def path_source_label(value: str) -> str:
    return str(value).strip() or "Importazione"
